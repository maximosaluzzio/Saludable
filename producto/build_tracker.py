"""Builds the Rental Property Tracker spreadsheet (product for Etsy/Gumroad)."""
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "Rental_Property_Tracker.xlsx"
YEAR = 2026
LOG_ROWS = 1000          # rows available in Income/Expenses logs
FIRST = 5                # first data row in logs
LAST = FIRST + LOG_ROWS - 1
MAX_PROPS = 10

# ---------- styles ----------
DARK = "1F4E5A"
ACCENT = "2E8B7A"
LIGHT = "E8F3F1"
GRAY = "F2F2F2"
INPUT = "FFF9E0"
F = "Arial"

title_font = Font(name=F, size=18, bold=True, color=DARK)
sub_font = Font(name=F, size=10, italic=True, color="666666")
hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=DARK)
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
calc_fill = PatternFill("solid", fgColor=GRAY)
input_fill = PatternFill("solid", fgColor=INPUT)
light_fill = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="D0D0D0")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
wrap = Alignment(wrap_text=True, vertical="top")

USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "mmm d, yyyy"

EXPENSE_CATS = [
    ("Advertising", 5), ("Auto and travel", 6), ("Cleaning and maintenance", 7),
    ("Commissions", 8), ("Insurance", 9), ("Legal and other professional fees", 10),
    ("Management fees", 11), ("Mortgage interest paid to banks", 12),
    ("Other interest", 13), ("Repairs", 14), ("Supplies", 15), ("Taxes", 16),
    ("Utilities", 17), ("Other", 19),
]
INCOME_TYPES = ["Rent", "Late fee", "Pet fee", "Parking", "Laundry / vending", "Other"]
PAY_METHODS = ["Bank transfer", "Check", "Cash", "Zelle", "Venmo", "Credit card", "Other"]


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, box
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = title_font
    ws["A2"] = subtitle
    ws["A2"].font = sub_font
    ws.row_dimensions[1].height = 30
    ws.sheet_view.showGridLines = False


def style_range(ws, rng, font=body, fill=None, fmt=None, border=True, align=None):
    for row in ws[rng]:
        for c in row:
            c.font = font
            if fill:
                c.fill = fill
            if fmt:
                c.number_format = fmt
            if border:
                c.border = box
            if align:
                c.alignment = align


wb = Workbook()

# ================= START HERE =================
ws = wb.active
ws.title = "Start Here"
title(ws, "Rental Property Tracker", "Income, expenses, rent roll & Schedule E summary — for small landlords")
ws["B1"], ws["B2"] = ws["A1"].value, ws["A2"].value
ws["B1"].font, ws["B2"].font = title_font, sub_font
ws["A1"], ws["A2"] = None, None
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 95
steps = [
    ("HOW TO USE", None),
    ("1.  Go to the Properties tab and type your properties (up to 10). The names you use there feed every dropdown.", None),
    ("2.  Log every payment you receive in the Income tab (date, property, type, amount).", None),
    ("3.  Log every bill you pay in the Expenses tab. Categories match IRS Schedule E lines, so tax time is copy-and-paste.", None),
    ("4.  Keep your tenants and lease dates in the Rent Roll tab — leases ending within 60 days light up automatically.", None),
    ("5.  Open the Dashboard, pick a property (or All Properties) and a year. Everything updates on its own.", None),
    ("6.  At tax time, open Schedule E Summary: totals per property, per IRS line, ready for you or your accountant.", None),
    ("", None),
    ("COLOR LEGEND", None),
    ("Light yellow cells = type here (inputs).", input_fill),
    ("Gray cells = automatic formulas. Don't type over them.", calc_fill),
    ("", None),
    ("BEFORE YOU START", None),
    ("The Income, Expenses, Properties and Rent Roll tabs contain sample rows so you can see how it works. "
     "Select those rows and press Delete (clear contents — don't delete the rows) to start fresh.", None),
    ("Works in Microsoft Excel (2010+), Google Sheets (File > Import) and Apple Numbers. No macros.", None),
    ("", None),
    ("This template is an organizational tool, not tax or legal advice. Confirm figures with a tax professional.", None),
]
r = 4
for text, fill in steps:
    c = ws.cell(row=r, column=2, value=text)
    c.alignment = wrap
    if text.isupper():
        c.font = Font(name=F, size=11, bold=True, color=ACCENT)
    else:
        c.font = body
    if fill:
        c.fill = fill
        c.border = box
    ws.row_dimensions[r].height = 30 if len(text) > 95 else 18
    r += 1

# ================= LISTS (hidden helper) =================
lists = wb.create_sheet("Lists")
lists["A1"], lists["B1"], lists["C1"], lists["D1"], lists["E1"] = (
    "Property picker", "Expense categories", "Schedule E line", "Income types", "Payment methods")
for c in lists[1]:
    c.font = bold
lists["A2"] = "All Properties"
for i in range(MAX_PROPS):
    lists.cell(row=3 + i, column=1, value=f'=IF(Properties!A{5 + i}="","",Properties!A{5 + i})')
for i, (cat, line) in enumerate(EXPENSE_CATS):
    lists.cell(row=2 + i, column=2, value=cat)
    lists.cell(row=2 + i, column=3, value=line)
for i, t in enumerate(INCOME_TYPES):
    lists.cell(row=2 + i, column=4, value=t)
for i, t in enumerate(PAY_METHODS):
    lists.cell(row=2 + i, column=5, value=t)
lists.sheet_state = "hidden"

PROP_LIST = f"=Properties!$A$5:$A${4 + MAX_PROPS}"
CAT_LIST = f"=Lists!$B$2:$B${1 + len(EXPENSE_CATS)}"
INC_LIST = f"=Lists!$D$2:$D${1 + len(INCOME_TYPES)}"
PAY_LIST = f"=Lists!$E$2:$E${1 + len(PAY_METHODS)}"
PICK_LIST = f"=Lists!$A$2:$A${2 + MAX_PROPS}"


def dv(ws, formula, rng, prompt=None):
    v = DataValidation(type="list", formula1=formula, allow_blank=True)
    v.error = "Pick a value from the list."
    v.errorTitle = "Invalid entry"
    if prompt:
        v.prompt, v.showInputMessage = prompt, True
    ws.add_data_validation(v)
    v.add(rng)


# ================= PROPERTIES =================
ws = wb.create_sheet("Properties")
title(ws, "Properties", "One row per property. The name in column A is what you'll pick everywhere else.")
header(ws, 4, ["Property name", "Address", "Type", "Units", "Purchase date", "Purchase price",
               "Cash invested\n(down + closing + rehab)", "Current loan balance", "Est. market value",
               "Equity", "Notes"],
       [22, 34, 16, 8, 14, 15, 18, 17, 16, 15, 30])
sample_props = [
    ("Maple Duplex", "123 Maple St, Austin, TX", "Duplex", 2, date(2021, 6, 15), 285000, 68000, 205000, 360000),
    ("Oak Condo", "45 Oak Ave #3B, Tampa, FL", "Condo", 1, date(2023, 3, 1), 210000, 52000, 160000, 225000),
]
for i in range(MAX_PROPS):
    rr = 5 + i
    style_range(ws, f"A{rr}:I{rr}", fill=input_fill)
    style_range(ws, f"K{rr}:K{rr}", fill=input_fill)
    ws[f"E{rr}"].number_format = DATE
    for col in "FGHI":
        ws[f"{col}{rr}"].number_format = USD0
    ws[f"J{rr}"] = f'=IF(I{rr}="","",I{rr}-H{rr})'
    style_range(ws, f"J{rr}:J{rr}", fill=calc_fill, fmt=USD0)
for i, row in enumerate(sample_props):
    for j, val in enumerate(row):
        ws.cell(row=5 + i, column=1 + j, value=val)
ws.freeze_panes = "B5"

# ================= INCOME =================
ws = wb.create_sheet("Income")
title(ws, "Income Log", "Every payment received. Pick property, type and method from the dropdowns.")
header(ws, 4, ["Date", "Property", "Unit / Tenant", "Type", "Amount", "Payment method", "Notes", "Month"],
       [14, 22, 22, 18, 14, 16, 34, 11])
style_range(ws, f"A{FIRST}:G{LAST}", fill=input_fill)
for rr in range(FIRST, LAST + 1):
    ws[f"A{rr}"].number_format = DATE
    ws[f"E{rr}"].number_format = USD
    ws[f"H{rr}"] = f'=IF(A{rr}="","",TEXT(A{rr},"yyyy-mm"))'
style_range(ws, f"H{FIRST}:H{LAST}", fill=calc_fill, align=Alignment(horizontal="center"))
dv(ws, PROP_LIST, f"B{FIRST}:B{LAST}")
dv(ws, INC_LIST, f"D{FIRST}:D{LAST}")
dv(ws, PAY_LIST, f"F{FIRST}:F{LAST}")
dd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
dd.error = "Enter a valid date, e.g. 1/15/2026"
ws.add_data_validation(dd)
dd.add(f"A{FIRST}:A{LAST}")
sample_income = []
for m in range(1, 4):
    sample_income += [
        (date(YEAR, m, 1), "Maple Duplex", "Unit A - J. Rivera", "Rent", 1450, "Zelle", ""),
        (date(YEAR, m, 3), "Maple Duplex", "Unit B - K. Chen", "Rent", 1395, "Bank transfer", ""),
        (date(YEAR, m, 1), "Oak Condo", "M. Patel", "Rent", 1850, "Bank transfer", ""),
    ]
sample_income += [
    (date(YEAR, 2, 8), "Maple Duplex", "Unit B - K. Chen", "Late fee", 75, "Zelle", "Paid 5 days late"),
    (date(YEAR, 1, 1), "Oak Condo", "M. Patel", "Parking", 50, "Bank transfer", "Covered spot"),
]
for i, row in enumerate(sample_income):
    for j, val in enumerate(row):
        ws.cell(row=FIRST + i, column=1 + j, value=val)
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:H{LAST}"

# ================= EXPENSES =================
ws = wb.create_sheet("Expenses")
title(ws, "Expense Log", "Every bill paid. Categories follow IRS Schedule E lines.")
header(ws, 4, ["Date", "Property", "Category (Schedule E)", "Vendor / Payee", "Description", "Amount",
               "Payment method", "Receipt saved?", "Sch. E line", "Month"],
       [14, 22, 30, 22, 32, 14, 16, 12, 10, 11])
style_range(ws, f"A{FIRST}:H{LAST}", fill=input_fill)
for rr in range(FIRST, LAST + 1):
    ws[f"A{rr}"].number_format = DATE
    ws[f"F{rr}"].number_format = USD
    ws[f"I{rr}"] = f'=IF(C{rr}="","",IFERROR(INDEX(Lists!$C$2:$C$15,MATCH(C{rr},Lists!$B$2:$B$15,0)),""))'
    ws[f"J{rr}"] = f'=IF(A{rr}="","",TEXT(A{rr},"yyyy-mm"))'
style_range(ws, f"I{FIRST}:J{LAST}", fill=calc_fill, align=Alignment(horizontal="center"))
dv(ws, PROP_LIST, f"B{FIRST}:B{LAST}")
dv(ws, CAT_LIST, f"C{FIRST}:C{LAST}")
dv(ws, PAY_LIST, f"G{FIRST}:G{LAST}")
dv(ws, '"Yes,No"', f"H{FIRST}:H{LAST}")
dd2 = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
ws.add_data_validation(dd2)
dd2.add(f"A{FIRST}:A{LAST}")
sample_exp = [
    (date(YEAR, 1, 5), "Maple Duplex", "Mortgage interest paid to banks", "Chase", "Jan mortgage - interest portion", 1012.40, "Bank transfer", "Yes"),
    (date(YEAR, 2, 5), "Maple Duplex", "Mortgage interest paid to banks", "Chase", "Feb mortgage - interest portion", 1009.85, "Bank transfer", "Yes"),
    (date(YEAR, 3, 5), "Maple Duplex", "Mortgage interest paid to banks", "Chase", "Mar mortgage - interest portion", 1007.28, "Bank transfer", "Yes"),
    (date(YEAR, 1, 10), "Maple Duplex", "Insurance", "State Farm", "Annual landlord policy", 1380, "Credit card", "Yes"),
    (date(YEAR, 2, 14), "Maple Duplex", "Repairs", "ABC Plumbing", "Unit B water heater valve", 285, "Credit card", "Yes"),
    (date(YEAR, 1, 20), "Oak Condo", "Other", "Oak Tower HOA", "HOA dues - Q1", 960, "Check", "Yes"),
    (date(YEAR, 1, 31), "Oak Condo", "Taxes", "Hillsborough County", "Property tax installment", 1150, "Bank transfer", "Yes"),
    (date(YEAR, 3, 2), "Oak Condo", "Cleaning and maintenance", "Sparkle Co.", "HVAC filter + service", 140, "Venmo", "No"),
    (date(YEAR, 3, 18), "Maple Duplex", "Utilities", "City of Austin", "Water - common meter", 118.60, "Bank transfer", "Yes"),
    (date(YEAR, 2, 1), "Oak Condo", "Advertising", "Zillow", "Listing boost", 45, "Credit card", "Yes"),
]
for i, row in enumerate(sample_exp):
    for j, val in enumerate(row):
        ws.cell(row=FIRST + i, column=1 + j, value=val)
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:J{LAST}"

# ================= RENT ROLL =================
ws = wb.create_sheet("Rent Roll")
title(ws, "Rent Roll & Leases", "Current tenants. Days left and status calculate automatically from today's date.")
header(ws, 4, ["Property", "Unit", "Tenant name", "Phone", "Email", "Lease start", "Lease end",
               "Monthly rent", "Security deposit held", "Days to lease end", "Status"],
       [22, 10, 22, 15, 26, 13, 13, 14, 15, 13, 18])
RR_LAST = 54
style_range(ws, f"A5:I{RR_LAST}", fill=input_fill)
for rr in range(5, RR_LAST + 1):
    ws[f"F{rr}"].number_format = DATE
    ws[f"G{rr}"].number_format = DATE
    ws[f"H{rr}"].number_format = USD0
    ws[f"I{rr}"].number_format = USD0
    ws[f"J{rr}"] = f'=IF(G{rr}="","",G{rr}-TODAY())'
    ws[f"K{rr}"] = (f'=IF(G{rr}="","",IF(J{rr}<0,"Expired / month-to-month",'
                    f'IF(J{rr}<=60,"Renew soon","Active")))')
style_range(ws, f"J5:K{RR_LAST}", fill=calc_fill, align=Alignment(horizontal="center"))
for rr in range(5, RR_LAST + 1):
    ws[f"J{rr}"].number_format = '0;-0;0'
dv(ws, PROP_LIST, f"A5:A{RR_LAST}")
red = PatternFill("solid", fgColor="F8D7DA")
amber = PatternFill("solid", fgColor="FFE8B3")
ws.conditional_formatting.add(f"J5:K{RR_LAST}", FormulaRule(formula=[f'AND($J5<>"",$J5<0)'], fill=red))
ws.conditional_formatting.add(f"J5:K{RR_LAST}", FormulaRule(formula=[f'AND($J5<>"",$J5>=0,$J5<=60)'], fill=amber))
sample_rr = [
    ("Maple Duplex", "A", "Jordan Rivera", "(512) 555-0142", "jrivera@email.com", date(2025, 8, 1), date(2026, 7, 31), 1450, 1450),
    ("Maple Duplex", "B", "Kim Chen", "(512) 555-0199", "kchen@email.com", date(2025, 12, 1), date(2026, 11, 30), 1395, 1395),
    ("Oak Condo", "3B", "Maya Patel", "(813) 555-0110", "mpatel@email.com", date(2026, 3, 1), date(2027, 2, 28), 1850, 2000),
]
for i, row in enumerate(sample_rr):
    for j, val in enumerate(row):
        ws.cell(row=5 + i, column=1 + j, value=val)
ws[f"A{RR_LAST + 2}"] = "Total monthly rent (scheduled)"
ws[f"A{RR_LAST + 2}"].font = bold
ws[f"H{RR_LAST + 2}"] = f"=SUM(H5:H{RR_LAST})"
ws[f"I{RR_LAST + 2}"] = f"=SUM(I5:I{RR_LAST})"
style_range(ws, f"H{RR_LAST + 2}:I{RR_LAST + 2}", font=bold, fill=calc_fill, fmt=USD0)
ws.freeze_panes = "A5"

# ================= DASHBOARD =================
ws = wb.create_sheet("Dashboard", 1)
title(ws, "Dashboard", "Pick a property and a year (yellow cells). Everything below updates automatically.")
for col, w in zip("ABCDEFGHIJ", [16, 15, 15, 15, 15, 3, 32, 15, 11, 3]):
    ws.column_dimensions[col].width = w
ws["A4"], ws["A5"] = "Property", "Year"
ws["A4"].font = ws["A5"].font = bold
ws["B4"], ws["B5"] = "All Properties", YEAR
ws.merge_cells("B4:C4")
for c in ("B4", "B5"):
    ws[c].fill, ws[c].font, ws[c].border = input_fill, Font(name=F, size=11, bold=True), box
ws["C4"].border = box
dv(ws, PICK_LIST, "B4")
ws["B5"].number_format = "0"
yv = DataValidation(type="whole", operator="between", formula1="2000", formula2="2100")
ws.add_data_validation(yv)
yv.add("B5")
# helper: wildcard when "All Properties"
ws["J4"] = '=IF(OR($B$4="All Properties",$B$4=""),"*",$B$4)'
ws["J4"].font = Font(name=F, size=8, color="FFFFFF")
P = "Dashboard!$J$4"
Y0 = "DATE($B$5,1,1)"
Y1 = "DATE($B$5+1,1,1)"
INC = f"Income!$E${FIRST}:$E${LAST}"
INC_D = f"Income!$A${FIRST}:$A${LAST}"
INC_P = f"Income!$B${FIRST}:$B${LAST}"
EXP = f"Expenses!$F${FIRST}:$F${LAST}"
EXP_D = f"Expenses!$A${FIRST}:$A${LAST}"
EXP_P = f"Expenses!$B${FIRST}:$B${LAST}"
EXP_C = f"Expenses!$C${FIRST}:$C${LAST}"

# KPI tiles
kpis = [
    ("A", "Total income", "=B23"),
    ("B", "Total expenses", "=C23"),
    ("C", "Net cash flow", "=D23"),
    ("D", "Expense ratio", "=IFERROR(C23/B23,0)"),
    ("E", "Avg. monthly net", "=IFERROR(D23/MAX(1,COUNTIF(D10:D21,\"<>0\")),0)"),
]
for col, label, f in kpis:
    ws[f"{col}7"] = label
    ws[f"{col}7"].font = Font(name=F, size=9, bold=True, color="FFFFFF")
    ws[f"{col}7"].fill = PatternFill("solid", fgColor=ACCENT)
    ws[f"{col}7"].alignment = center
    ws[f"{col}8"] = f
    ws[f"{col}8"].font = Font(name=F, size=14, bold=True, color=DARK)
    ws[f"{col}8"].fill = light_fill
    ws[f"{col}8"].alignment = center
    ws[f"{col}8"].number_format = PCT if label == "Expense ratio" else USD0
ws.row_dimensions[8].height = 32

# Monthly table
for i, lab in enumerate(["Month", "Income", "Expenses", "Net", "Cumulative net"]):
    c = ws.cell(row=9, column=1 + i, value=lab)
    c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, box
for m in range(1, 13):
    rr = 9 + m
    s, e = f"DATE($B$5,{m},1)", f"DATE($B$5,{m + 1},1)"
    ws[f"A{rr}"] = f"=DATE($B$5,{m},1)"
    ws[f"B{rr}"] = f'=SUMIFS({INC},{INC_P},{P},{INC_D},">="&{s},{INC_D},"<"&{e})'
    ws[f"C{rr}"] = f'=SUMIFS({EXP},{EXP_P},{P},{EXP_D},">="&{s},{EXP_D},"<"&{e})'
    ws[f"D{rr}"] = f"=B{rr}-C{rr}"
    ws[f"E{rr}"] = f"=D{rr}" if m == 1 else f"=E{rr - 1}+D{rr}"
style_range(ws, "A10:E21", fill=calc_fill, fmt=USD0)
for rr in range(10, 22):
    ws[f"A{rr}"].number_format = "mmm"
    ws[f"A{rr}"].alignment = Alignment(horizontal="center")
ws.conditional_formatting.add("D10:E21", CellIsRule(operator="lessThan", formula=["0"],
                                                    font=Font(name=F, size=10, color="C00000")))
ws["A23"] = "Total"
ws["B23"], ws["C23"], ws["D23"] = "=SUM(B10:B21)", "=SUM(C10:C21)", "=B23-C23"
style_range(ws, "A23:D23", font=bold, fill=light_fill, fmt=USD0)

# Expenses by category
for i, lab in enumerate(["Expenses by category", "Amount", "% of total"]):
    c = ws.cell(row=9, column=7 + i, value=lab)
    c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, box
for i, (cat, _) in enumerate(EXPENSE_CATS):
    rr = 10 + i
    ws[f"G{rr}"] = cat
    ws[f"H{rr}"] = f'=SUMIFS({EXP},{EXP_P},{P},{EXP_C},G{rr},{EXP_D},">="&{Y0},{EXP_D},"<"&{Y1})'
    ws[f"I{rr}"] = f"=IFERROR(H{rr}/$C$23,0)"
cat_last = 9 + len(EXPENSE_CATS)
style_range(ws, f"G10:G{cat_last}")
style_range(ws, f"H10:H{cat_last}", fill=calc_fill, fmt=USD0)
style_range(ws, f"I10:I{cat_last}", fill=calc_fill, fmt=PCT)
ws.conditional_formatting.add(f"I10:I{cat_last}", DataBarRule(start_type="num", start_value=0,
                                                              end_type="num", end_value=1, color=ACCENT))

# Per-property summary
ps = 27
ws[f"A{ps - 1}"] = "Portfolio summary (selected year, all properties)"
ws[f"A{ps - 1}"].font = Font(name=F, size=12, bold=True, color=DARK)
for col, lab in zip("ABCDEH", ["Property", "Income", "Expenses", "Net cash flow", "Cash invested", "Equity"]):
    c = ws[f"{col}{ps}"]
    c.value = lab
    c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, box
ws[f"G{ps}"] = "Cash-on-cash return"
c = ws[f"G{ps}"]
c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, box
for i in range(MAX_PROPS):
    rr = ps + 1 + i
    pr = 5 + i
    ws[f"A{rr}"] = f'=IF(Properties!A{pr}="","",Properties!A{pr})'
    ws[f"B{rr}"] = (f'=IF(A{rr}="","",SUMIFS({INC},{INC_P},A{rr},{INC_D},">="&{Y0},{INC_D},"<"&{Y1}))')
    ws[f"C{rr}"] = (f'=IF(A{rr}="","",SUMIFS({EXP},{EXP_P},A{rr},{EXP_D},">="&{Y0},{EXP_D},"<"&{Y1}))')
    ws[f"D{rr}"] = f'=IF(A{rr}="","",B{rr}-C{rr})'
    ws[f"E{rr}"] = f'=IF(A{rr}="","",Properties!G{pr})'
    ws[f"G{rr}"] = f'=IF(OR(A{rr}="",N(E{rr})=0),"",D{rr}/E{rr})'
    ws[f"H{rr}"] = f'=IF(A{rr}="","",Properties!J{pr})'
pl = ps + MAX_PROPS
style_range(ws, f"A{ps + 1}:A{pl}", fill=calc_fill)
style_range(ws, f"B{ps + 1}:E{pl}", fill=calc_fill, fmt=USD0)
style_range(ws, f"G{ps + 1}:G{pl}", fill=calc_fill, fmt=PCT)
style_range(ws, f"H{ps + 1}:H{pl}", fill=calc_fill, fmt=USD0)
ws[f"G{ps}"].comment = Comment(
    "Net cash flow for the selected year divided by cash invested (Properties tab). "
    "If the year is still in progress, this is year-to-date, not annualized.", "Template")

# Chart
ch = BarChart()
ch.type = "col"
ch.title = "Income vs. expenses by month"
ch.y_axis.title = None
ch.y_axis.numFmt = "$#,##0"
ch.y_axis.majorGridlines = None
ch.height, ch.width = 7.5, 17
data = Reference(ws, min_col=2, max_col=3, min_row=9, max_row=21)
cats = Reference(ws, min_col=1, min_row=10, max_row=21)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.x_axis.number_format = "mmm"
ch.series[0].graphicalProperties.solidFill = ACCENT
ch.series[0].graphicalProperties.line.solidFill = ACCENT
ch.series[1].graphicalProperties.solidFill = "E07A5F"
ch.series[1].graphicalProperties.line.solidFill = "E07A5F"
ch.legend.position = "b"
ws.add_chart(ch, f"A{pl + 3}")

# ================= SCHEDULE E =================
ws = wb.create_sheet("Schedule E Summary", 2)
title(ws, "Schedule E Summary", "Totals per IRS Schedule E (Form 1040) line, per property, for the year below.")
ws["A4"], ws["B4"] = "Tax year", YEAR
ws["A4"].font = bold
ws["B4"].fill, ws["B4"].font, ws["B4"].border = input_fill, Font(name=F, size=11, bold=True), box
ws["B4"].number_format = "0"
ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 34
hr = 6
ws.cell(row=hr, column=1, value="Line")
ws.cell(row=hr, column=2, value="Description")
for i in range(MAX_PROPS):
    col = 3 + i
    ws.cell(row=hr, column=col, value=f'=IF(Properties!A{5 + i}="","Property {i + 1}",Properties!A{5 + i})')
    ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 15
tot_col = 3 + MAX_PROPS
ws.cell(row=hr, column=tot_col, value="All properties")
ws.column_dimensions[ws.cell(row=1, column=tot_col).column_letter].width = 16
for c in ws[hr][:tot_col]:
    c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, center, box
ws.row_dimensions[hr].height = 32

SY0, SY1 = "DATE($B$4,1,1)", "DATE($B$4+1,1,1)"
rows = [("3", "Rents received", "income")] + \
       [(str(line), cat, "exp") for cat, line in EXPENSE_CATS if line < 19] + \
       [("18", "Depreciation expense or depletion (enter manually)", "manual"),
        ("19", "Other (HOA, pest control, bank fees, etc.)", "exp_other"),
        ("20", "Total expenses (lines 5-19)", "total"),
        ("21", "Income or (loss) (line 3 minus line 20)", "net")]
r = hr + 1
row_of = {}
for line, desc, kind in rows:
    ws.cell(row=r, column=1, value=line).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=2, value=desc)
    row_of[line] = r
    for i in range(MAX_PROPS):
        col = 3 + i
        cl = ws.cell(row=1, column=col).column_letter
        pcell = f"Properties!$A${5 + i}"
        if kind == "income":
            f = f'=IF({pcell}="",0,SUMIFS({INC},{INC_P},{pcell},{INC_D},">="&{SY0},{INC_D},"<"&{SY1}))'
        elif kind in ("exp", "exp_other"):
            cat = "Other" if kind == "exp_other" else desc
            f = (f'=IF({pcell}="",0,SUMIFS({EXP},{EXP_P},{pcell},{EXP_C},"{cat}",'
                 f'{EXP_D},">="&{SY0},{EXP_D},"<"&{SY1}))')
        elif kind == "manual":
            f = None
        elif kind == "total":
            f = f"=SUM({cl}{row_of['5']}:{cl}{r - 1})"
        else:
            f = f"={cl}{row_of['3']}-{cl}{row_of['20']}"
        c = ws.cell(row=r, column=col, value=f)
        c.number_format = USD
        c.border = box
        c.font = bold if kind in ("total", "net", "income") else body
        c.fill = input_fill if kind == "manual" else calc_fill
    tl = ws.cell(row=1, column=tot_col - 1).column_letter
    c = ws.cell(row=r, column=tot_col, value=f"=SUM(C{r}:{tl}{r})")
    c.number_format, c.border, c.fill = USD, box, light_fill
    c.font = bold
    for cc in (ws.cell(row=r, column=1), ws.cell(row=r, column=2)):
        cc.border = box
        cc.font = bold if kind in ("total", "net", "income") else body
    r += 1
ws.cell(row=row_of["18"], column=2).comment = Comment(
    "Depreciation isn't a cash expense, so it isn't in the Expense log. "
    "Typical residential rental: building basis / 27.5 years (IRS Pub. 527). Ask your tax pro.", "Template")
ws.cell(row=r + 1, column=2, value="Line 4 (royalties) and lines 22+ (passive loss limits) are not covered. "
        "Not tax advice — review with a tax professional.").font = sub_font
ws.freeze_panes = "C7"

# sheet tab colors / order
for name, color in [("Start Here", ACCENT), ("Dashboard", DARK), ("Schedule E Summary", DARK),
                    ("Properties", "C9A227"), ("Income", "C9A227"), ("Expenses", "C9A227"),
                    ("Rent Roll", "C9A227")]:
    wb[name].sheet_properties.tabColor = color
for w in wb.worksheets:
    w.sheet_view.zoomScale = 100
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
for name, last in [("Income", "H200"), ("Expenses", "J200"), ("Rent Roll", f"K{RR_LAST + 2}")]:
    wb[name].print_area = f"A1:{last}"
    wb[name].print_title_rows = "4:4"
wb.active = 0
wb.save(OUT)
print("saved", OUT)
